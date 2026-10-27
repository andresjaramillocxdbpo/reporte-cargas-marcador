from datetime import date, datetime
from decimal import Decimal

from openpyxl import load_workbook

from reporte_cargas import excel
from reporte_cargas.extraccion import _calcular_pct_estados

DIA = date(2026, 10, 23)


def _resumen(idload, carga, marcaciones, agregados=0):
    return {
        "fecha_marcacion": DIA, "marcador": "ENIGMA_CXD_1", "idload": idload, "carga": carga,
        "tipo_base": "HOGAR OUTBOUND", "skilldef": "hogarelitemotor",
        "fecha_creacion_carga": datetime(2026, 10, 21, 16, 51, 50),
        "registros_cti_loads": 39222, "registros_inicio_dia": 39222,
        "registros_cargados": 39222 + agregados, "registros_agregados_dia": agregados,
        "crecio_despues_de_creada": int(agregados > 0), "numeros_unicos": 39222,
        "marcaciones_dia": marcaciones, "numeros_marcados_dia": 5000,
        "min_intentos_dia": 1, "prom_intentos_dia": Decimal("1.56"), "max_intentos_dia": 4,
        "registros_marcados_hist": 39222, "registros_sin_marcar_hist": 0,
        "numeros_marcados_hist": 39222, "pct_avance_hist": Decimal("100.00"),
        "min_intentos_hist": 1, "prom_intentos_hist": Decimal("3.10"), "max_intentos_hist": 9,
    }


def _estado(idload, resultdesc, marcaciones):
    return {"fecha_marcacion": DIA, "marcador": "ENIGMA_CXD_1", "idload": idload,
            "resultdesc": resultdesc, "marcaciones": marcaciones, "numeros": marcaciones}


def test_pct_estados_suma_100_por_carga():
    estados = [_estado(461, "NO CONTESTA", 6000), _estado(461, "CONTESTA", 1807),
               _estado(462, "NO CONTESTA", 44296)]
    _calcular_pct_estados(estados)
    assert sum(e["pct_marcaciones"] for e in estados if e["idload"] == 461) == Decimal("100.00")
    assert estados[2]["pct_marcaciones"] == Decimal("100.00")


def test_excel_tiene_todas_las_hojas(tmp_path):
    resumen = [_resumen(461, "MOTOR ELITE - CURAZAO", 7807), _resumen(462, "OTRA", 44296, agregados=500)]
    estados = [_estado(461, "NO CONTESTA", 6000), _estado(461, "CONTESTA", 1807),
               _estado(462, "NO CONTESTA", 44296)]
    estados_hist = [_estado(461, "NO CONTESTA", 30000), _estado(461, "CONTESTA", 17000),
                    _estado(461, "BUZON", 100)]
    _calcular_pct_estados(estados)
    _calcular_pct_estados(estados_hist)

    ruta = excel.generar(resumen, estados, estados_hist, DIA, DIA, str(tmp_path))
    wb = load_workbook(ruta)

    assert wb.sheetnames == ["Resumen", "Crecimiento", "Estados día", "Estados histórico",
                             "Estados detalle", "Glosario"]
    # cada hoja con datos es una Tabla de Excel con nombre
    tablas = {n for ws in wb.worksheets for n in ws.tables}
    assert {"tbl_resumen", "tbl_crecimiento", "tbl_estados_dia", "tbl_estados_hist",
            "tbl_estados_detalle", "tbl_glosario"} <= tablas
    assert wb["Resumen"].max_row == 3
    assert wb["Crecimiento"].max_row == 2 and wb["Crecimiento"]["C2"].value == 462

    encabezados = [c.value for c in wb["Estados día"][1]]
    assert encabezados[-2:] == ["% NO CONTESTA", "% CONTESTA"]
    fila_461 = [c.value for c in wb["Estados día"][2]]
    assert fila_461[2] == 461 and fila_461[4] == 7807

    acum = [c.value for c in wb["Estados histórico"][1]]
    assert acum[0] == "Corte (fin del día)" and acum[-1] == "% BUZON"
    assert wb["Estados histórico"]["E2"].value == 47100

    detalle = wb["Estados detalle"]
    assert detalle.max_row == 1 + len(estados) + len(estados_hist)
    assert {detalle.cell(r, 1).value for r in range(2, detalle.max_row + 1)} == {"Día", "Histórico"}
    # el detalle trae los datos de la carga para filtrar sin cruzar hojas
    assert detalle["E2"].value == "MOTOR ELITE - CURAZAO" and detalle["F2"].value == "HOGAR OUTBOUND"


def test_excel_vacio_no_falla(tmp_path):
    ruta = excel.generar([], [], [], DIA, DIA, str(tmp_path))
    assert ruta.exists()
