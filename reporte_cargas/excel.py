"""Generación del Excel: Resumen, Crecimiento, Estados día, Estados histórico (pivots en %),
Estados detalle y Glosario. Cada hoja con datos es una Tabla de Excel con nombre (tbl_*)."""

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .modelo import COLUMNAS_ESTADOS, COLUMNAS_RESUMEN, DESCRIPCIONES

ENCABEZADO_FONDO = PatternFill("solid", fgColor="1F4E78")
ENCABEZADO_FUENTE = Font(bold=True, color="FFFFFF")

FORMATOS = {
    "fecha_marcacion": "yyyy-mm-dd",
    "fecha_creacion_carga": "yyyy-mm-dd hh:mm",
    "prom_intentos_dia": "0.00",
    "prom_intentos_hist": "0.00",
    "pct_avance_hist": "0.00",
    "pct_marcaciones": "0.00",
}
FORMATO_ENTERO = "#,##0"


def _escribir_hoja(ws, titulos: list[str], filas: list[list], formatos: list[str | None],
                   tabla: str | None = None) -> None:
    """Escribe encabezado y filas. Con `tabla`, el rango queda como Tabla de Excel con ese
    nombre (filtros, y lista para tablas dinámicas o Power Query)."""
    ws.append(titulos)
    for celda in ws[1]:
        celda.fill = ENCABEZADO_FONDO
        celda.font = ENCABEZADO_FUENTE
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32

    for fila in filas:
        ws.append(fila)

    for i, (titulo, formato) in enumerate(zip(titulos, formatos), start=1):
        letra = get_column_letter(i)
        if formato:
            for (celda,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                celda.number_format = formato
        ancho = max([len(str(titulo))] + [len(str(f[i - 1])) for f in filas if f[i - 1] is not None])
        ws.column_dimensions[letra].width = min(max(ancho + 2, 10), 50)

    ws.freeze_panes = "A2"
    if filas and tabla:
        t = Table(displayName=tabla, ref=ws.dimensions)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)
    elif filas:
        ws.auto_filter.ref = ws.dimensions


def _formato(col: str, valor_ejemplo) -> str | None:
    if col in FORMATOS:
        return FORMATOS[col]
    if isinstance(valor_ejemplo, int) and col != "idload":
        return FORMATO_ENTERO
    return None


def _ejemplo(filas: list[dict], col: str):
    return next((f[col] for f in filas if f[col] is not None), None)


def _hoja_estados(ws, estados: list[dict], nombres: dict, titulo_fecha: str, tabla: str) -> None:
    """Una fila por carga y día, una columna por resultado con su %."""
    totales: dict[str, int] = {}
    for e in estados:
        totales[e["resultdesc"]] = totales.get(e["resultdesc"], 0) + e["marcaciones"]
    resultados = sorted(totales, key=lambda r: -totales[r])

    pivot: dict[tuple, dict] = {}
    for e in estados:
        clave = (e["fecha_marcacion"], e["marcador"], e["idload"])
        fila = pivot.setdefault(clave, {"marcaciones": 0, "pct": {}})
        fila["marcaciones"] += e["marcaciones"]
        fila["pct"][e["resultdesc"]] = e["pct_marcaciones"]

    filas = [
        [clave[0], clave[1], clave[2], nombres.get(clave), datos["marcaciones"]]
        + [datos["pct"].get(r) for r in resultados]
        for clave, datos in sorted(pivot.items())
    ]
    _escribir_hoja(
        ws,
        [titulo_fecha, "Marcador", "ID carga", "Carga", "Marcaciones"] + [f"% {r}" for r in resultados],
        filas,
        ["yyyy-mm-dd", None, None, None, FORMATO_ENTERO] + ["0.00"] * len(resultados),
        tabla,
    )


def generar(resumen: list[dict], estados: list[dict], estados_hist: list[dict],
            desde: date, hasta: date, carpeta: str) -> Path:
    resumen = sorted(resumen, key=lambda f: (f["fecha_marcacion"], f["marcador"], f["tipo_base"] or "", f["idload"]))
    orden_estados = lambda f: (f["fecha_marcacion"], f["marcador"], f["idload"], -f["marcaciones"])  # noqa: E731
    estados = sorted(estados, key=orden_estados)
    estados_hist = sorted(estados_hist, key=orden_estados)

    wb = Workbook()

    # Resumen
    ws = wb.active
    ws.title = "Resumen"
    cols = [c for c, _ in COLUMNAS_RESUMEN]
    _escribir_hoja(
        ws,
        [t for _, t in COLUMNAS_RESUMEN],
        [[f[c] for c in cols] for f in resumen],
        [_formato(c, _ejemplo(resumen, c)) for c in cols],
        "tbl_resumen",
    )

    # Crecimiento: cargas que recibieron registros después del día en que se crearon
    ws = wb.create_sheet("Crecimiento")
    cols_crec = ["fecha_marcacion", "marcador", "idload", "carga", "tipo_base", "fecha_creacion_carga",
                 "registros_inicio_dia", "registros_cargados", "registros_agregados_dia"]
    titulos = dict(COLUMNAS_RESUMEN)
    crecieron = [f for f in resumen if f["crecio_despues_de_creada"]]
    _escribir_hoja(
        ws,
        [titulos[c] for c in cols_crec],
        [[f[c] for c in cols_crec] for f in crecieron],
        [_formato(c, _ejemplo(resumen, c)) for c in cols_crec],
        "tbl_crecimiento",
    )

    nombres = {(f["fecha_marcacion"], f["marcador"], f["idload"]): f["carga"] for f in resumen}
    _hoja_estados(wb.create_sheet("Estados día"), estados, nombres, "Fecha marcación", "tbl_estados_dia")
    _hoja_estados(wb.create_sheet("Estados histórico"), estados_hist, nombres, "Corte (fin del día)",
                  "tbl_estados_hist")

    # Estados detalle: formato largo con los datos de la carga, para filtrar y
    # armar tablas dinámicas sin cruzar hojas
    ws = wb.create_sheet("Estados detalle")
    info = {(f["fecha_marcacion"], f["marcador"], f["idload"]): f for f in resumen}
    detalle = [("Día", f) for f in estados] + [("Histórico", f) for f in estados_hist]
    filas = []
    for alcance, e in detalle:
        carga = info.get((e["fecha_marcacion"], e["marcador"], e["idload"]), {})
        filas.append([alcance, e["fecha_marcacion"], e["marcador"], e["idload"], carga.get("carga"),
                      carga.get("tipo_base"), carga.get("skilldef"), e["resultdesc"],
                      e["marcaciones"], e["numeros"], e["pct_marcaciones"]])
    _escribir_hoja(
        ws,
        ["Alcance", "Fecha marcación", "Marcador", "ID carga", "Carga", "Tipo base", "Skill",
         "Resultado", "Marcaciones", "Números", "% marcaciones"],
        filas,
        [None, "yyyy-mm-dd", None, None, None, None, None, None, FORMATO_ENTERO, FORMATO_ENTERO, "0.00"],
        "tbl_estados_detalle",
    )

    # Glosario: qué significa cada columna
    ws = wb.create_sheet("Glosario")
    titulos = dict(COLUMNAS_RESUMEN) | dict(COLUMNAS_ESTADOS)
    _escribir_hoja(
        ws,
        ["Columna", "Nombre técnico", "Descripción"],
        [[titulos.get(c, c), c, d] for c, d in DESCRIPCIONES.items()]
        + [["Alcance", "-", "En Estados detalle: Día = solo el día reportado; Histórico = desde la creación "
                            "de la carga hasta el fin del día reportado."]],
        [None, None, None],
        "tbl_glosario",
    )
    ws.column_dimensions["C"].width = 90

    Path(carpeta).mkdir(parents=True, exist_ok=True)
    sufijo = desde.isoformat() if desde == hasta else f"{desde.isoformat()}_a_{hasta.isoformat()}"
    ruta = Path(carpeta) / f"reporte_cargas_{sufijo}.xlsx"
    try:
        wb.save(ruta)
    except PermissionError:
        # El archivo está abierto en Excel: se guarda con la hora para no perder la corrida
        ruta = ruta.with_stem(f"{ruta.stem}_{datetime.now():%H%M%S}")
        wb.save(ruta)
    return ruta
